from __future__ import annotations

import csv
import hashlib
import json
import re
import unicodedata
from collections import defaultdict
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .accounting_report_proof import PRESENTATION_CSV


REQUIRED_OFFICE_CASES = {
    ("trial-balance", "pre-closing", "1"),
    ("trial-balance", "pre-closing", "2"),
    ("trial-balance", "post-closing", "1"),
    ("trial-balance", "post-closing", "2"),
    ("balance-sheet", "post-closing", "1"),
    ("balance-sheet", "post-closing", "2"),
    ("income-statement", "pre-closing", "1"),
    ("income-statement", "pre-closing", "2"),
}
CLASSIFIERS = {
    "ACTIVO": "001",
    "PASIVO": "002",
    "CAPITAL": "003",
    "RESULTADO DEUDORAS": "004",
    "RESULTADO ACREEDORAS": "005",
}
AMOUNT_FIELDS = ("closing_balance", "accumulated_balance")


def _decimal(value: Any) -> Decimal:
    return Decimal(str(value or 0)).quantize(Decimal("0.01"))


def _json_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, dict):
        return {key: _json_value(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_value(item) for item in value]
    return value


def _content_sha256(value: Any) -> str:
    payload = json.dumps(_json_value(value), sort_keys=True, separators=(",", ":"), ensure_ascii=False)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _normalize_label(value: Any) -> str:
    decomposed = unicodedata.normalize("NFKD", str(value or ""))
    without_marks = "".join(character for character in decomposed if not unicodedata.combining(character))
    return " ".join(without_marks.upper().split())


def _optional_int(value: str | None) -> int | None:
    return int(value) if value not in (None, "") else None


def _load_presentation() -> list[dict[str, Any]]:
    with PRESENTATION_CSV.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    for row in rows:
        for field in ("trial_balance_order", "income_statement_order", "account_order"):
            row[field] = _optional_int(row[field])
        for field in ("trial_balance_selected", "balance_sheet_selected"):
            row[field] = row[field] == "1"
    return rows


def _selected(rows: list[dict[str, Any]], report_type: str) -> list[dict[str, Any]]:
    if report_type == "trial-balance":
        selected = [row for row in rows if row["trial_balance_selected"]]
    elif report_type == "balance-sheet":
        selected = [row for row in rows if row["balance_sheet_selected"]]
    else:
        selected = [
            row
            for row in rows
            if row["trial_balance_selected"]
            and not row["balance_sheet_selected"]
            and row["classifier_code"] in {"004", "005"}
        ]
    if report_type == "income-statement":
        return sorted(
            selected,
            key=lambda row: (
                1 if row["classifier_code"] == "005" else 2,
                row["account_order"] if row["account_order"] is not None else 9223372036854775807,
                row["gl_code"],
            ),
        )
    return sorted(
        selected,
        key=lambda row: (
            row["trial_balance_order"]
            if row["trial_balance_order"] is not None
            else row["income_statement_order"] or 32767,
            row["account_order"] if row["account_order"] is not None else 9223372036854775807,
            row["gl_code"],
        ),
    )


def _file_identity(path: Path) -> tuple[str, str, str]:
    name = _normalize_label(path.name)
    if name.startswith("BALANCE-COMPROBACION-PRE-LIQUIDACION"):
        report_type, closing_mode = "trial-balance", "pre-closing"
    elif name.startswith("BALANCE-COMPROBACION"):
        report_type, closing_mode = "trial-balance", "post-closing"
    elif name.startswith("BALANCE-GRAL"):
        report_type, closing_mode = "balance-sheet", "post-closing"
    elif name.startswith("ESTADO DE RESULTADOS-PRE-LIQ"):
        report_type, closing_mode = "income-statement", "pre-closing"
    elif name.startswith("ESTADO DE RESULTADOS"):
        report_type, closing_mode = "income-statement", "post-closing"
    else:
        raise ValueError("unsupported authoritative report filename")
    if "SANTIAGO" in name:
        scope = "1"
    elif "USULUTAN" in name:
        scope = "2"
    else:
        raise ValueError("authoritative report filename does not identify Santiago or Usulutan")
    return report_type, closing_mode, scope


def _workbook_rows(path: Path) -> list[tuple[Any, ...]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return [tuple(row) for sheet in workbook.worksheets for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _header_scope(rows: list[tuple[Any, ...]]) -> str | None:
    header = " ".join(
        _normalize_label(value)
        for row in rows
        for value in row
        if isinstance(value, str)
        and ("SUCURSAL:" in _normalize_label(value) or " - AGENCIA CENTRAL" in _normalize_label(value) or " - USULUTAN" in _normalize_label(value))
    )
    if "USULUTAN" in header:
        return "2"
    if "AGENCIA CENTRAL" in header:
        return "1"
    return None


def _generation_timestamp(rows: list[tuple[Any, ...]]) -> str | None:
    pattern = re.compile(r"\b\d{2}/\d{2}/\d{4}\s+\d{2}:\d{2}\b")
    for row in reversed(rows):
        for value in row:
            match = pattern.search(str(value or ""))
            if match:
                return match.group(0)
    return None


def _report_rows(
    workbook_rows: list[tuple[Any, ...]], report_type: str, presentation: list[dict[str, Any]]
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    selected = _selected(presentation, report_type)
    selected_by_code = {row["gl_code"]: row for row in selected}
    label_index: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in selected:
        label_index[(row["classifier_code"], _normalize_label(row["source_name"]))].append(row)

    observed: dict[str, dict[str, Decimal]] = {}
    observed_order: list[str] = []
    findings: list[dict[str, Any]] = []
    classifier_code: str | None = None
    for row_number, values in enumerate(workbook_rows, start=1):
        first = _normalize_label(values[0] if len(values) > 0 else None)
        if first in CLASSIFIERS:
            classifier_code = CLASSIFIERS[first]
        label = values[1] if len(values) > 1 else None
        if report_type == "income-statement":
            match = re.match(r"^(\d{4})\s+(.+)$", str(label or "").strip())
            if not match:
                continue
            code = match.group(1)
            if code not in selected_by_code:
                findings.append({"row": row_number, "reason": "UNMAPPED_GL_CODE", "gl_code": code})
                continue
            monthly = values[6] if len(values) > 6 else None
            accumulated = values[9] if len(values) > 9 else None
            if not isinstance(monthly, (int, float, Decimal)) or not isinstance(accumulated, (int, float, Decimal)):
                findings.append({"row": row_number, "reason": "MISSING_INCOME_AMOUNT", "gl_code": code})
                continue
            observed[code] = {
                "closing_balance": _decimal(monthly),
                "accumulated_balance": _decimal(accumulated),
            }
            observed_order.append(code)
            continue
        amount = values[8] if len(values) > 8 else None
        if not label or not isinstance(amount, (int, float, Decimal)):
            continue
        candidates = label_index.get((classifier_code or "", _normalize_label(label)), [])
        if len(candidates) != 1:
            findings.append(
                {
                    "row": row_number,
                    "reason": "UNMAPPED_OR_AMBIGUOUS_LABEL",
                    "label": str(label),
                    "classifier_code": classifier_code,
                    "candidate_gl_codes": [candidate["gl_code"] for candidate in candidates],
                }
            )
            continue
        code = candidates[0]["gl_code"]
        observed[code] = {"closing_balance": _decimal(amount)}
        observed_order.append(code)

    selected_order = [row["gl_code"] for row in selected]
    selected_positions = {code: index for index, code in enumerate(selected_order)}
    if observed_order != sorted(observed_order, key=selected_positions.__getitem__):
        findings.append({"reason": "VISIBLE_ROW_ORDER_MISMATCH", "observed_gl_codes": observed_order})

    rows = []
    for item in selected:
        row = {"gl_code": item["gl_code"], "account_name": item["source_name"]}
        row.update(observed.get(item["gl_code"], {"closing_balance": Decimal()}))
        if report_type == "income-statement" and "accumulated_balance" not in row:
            row["accumulated_balance"] = Decimal()
        rows.append(row)
    return rows, findings


def _presentation_controls(rows: list[dict[str, Any]], presentation: list[dict[str, Any]], report_type: str) -> dict[str, Decimal]:
    metadata = {row["gl_code"]: row for row in presentation}

    def total(field: str, classifiers: set[str]) -> Decimal:
        return sum(
            (_decimal(row.get(field)) for row in rows if metadata[row["gl_code"]]["classifier_code"] in classifiers),
            Decimal(),
        ).quantize(Decimal("0.01"))

    if report_type == "income-statement":
        monthly_income = total("closing_balance", {"005"})
        monthly_expense = total("closing_balance", {"004"})
        accumulated_income = total("accumulated_balance", {"005"})
        accumulated_expense = total("accumulated_balance", {"004"})
        return {
            "monthly_income": monthly_income,
            "monthly_expense": monthly_expense,
            "monthly_result": monthly_income - monthly_expense,
            "accumulated_income": accumulated_income,
            "accumulated_expense": accumulated_expense,
            "accumulated_result": accumulated_income - accumulated_expense,
        }
    assets = total("closing_balance", {"001"})
    liabilities = total("closing_balance", {"002"})
    equity = total("closing_balance", {"003"})
    income = total("closing_balance", {"005"})
    expense = total("closing_balance", {"004"})
    controls = {
        "assets": assets,
        "liabilities": liabilities,
        "equity": equity,
        "income": income,
        "expense": expense,
    }
    if report_type == "trial-balance":
        controls["left_total"] = assets + expense
        controls["right_total"] = liabilities + equity + income
        controls["presentation_variance"] = controls["left_total"] - controls["right_total"]
    else:
        controls["current_result"] = Decimal()
        controls["liabilities_equity"] = liabilities + equity
        controls["accounting_equation_variance"] = assets - controls["liabilities_equity"]
    return controls


def _derived_consolidated(cases: dict[tuple[str, str, str], dict[str, Any]], presentation: list[dict[str, Any]]) -> list[dict[str, Any]]:
    consolidated = []
    for report_type, closing_mode in sorted({(key[0], key[1]) for key in REQUIRED_OFFICE_CASES}):
        left = cases[(report_type, closing_mode, "1")]
        right = cases[(report_type, closing_mode, "2")]
        rows = []
        for left_row, right_row in zip(left["rows"], right["rows"], strict=True):
            if left_row["gl_code"] != right_row["gl_code"] or left_row["account_name"] != right_row["account_name"]:
                raise ValueError("office report presentation rows do not align")
            row = {"gl_code": left_row["gl_code"], "account_name": left_row["account_name"]}
            for field in AMOUNT_FIELDS:
                if field in left_row or field in right_row:
                    row[field] = _decimal(left_row.get(field)) + _decimal(right_row.get(field))
            rows.append(row)
        consolidated.append(
            {
                "report_type": report_type,
                "closing_mode": closing_mode,
                "scope": "consolidated",
                "scope_derivation": "office-1-plus-office-2",
                "zero_row_policy": "source-omitted-rows-filled-from-frozen-presentation",
                "rows": rows,
                "controls": _presentation_controls(rows, presentation, report_type),
            }
        )
    return consolidated


def normalize_authoritative_report_directory(input_dir: Path) -> dict[str, Any]:
    presentation = _load_presentation()
    blockers: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    cases: dict[tuple[str, str, str], dict[str, Any]] = {}
    exports = []
    paths = sorted(input_dir.glob("*.xlsx"))
    if not paths:
        blockers.append({"reason": "NO_XLSX_EXPORTS", "input_dir": str(input_dir.resolve())})
    for path in paths:
        try:
            report_type, closing_mode, filename_scope = _file_identity(path)
        except ValueError as exc:
            warnings.append({"file": path.name, "reason": "UNSUPPORTED_EXPORT", "detail": str(exc)})
            continue
        rows = _workbook_rows(path)
        header_scope = _header_scope(rows)
        required = (report_type, closing_mode, filename_scope) in REQUIRED_OFFICE_CASES
        identity_ok = header_scope == filename_scope
        date_ok = any("31 DE DICIEMBRE DE 2024" in _normalize_label(value) for row in rows for value in row)
        normalized_rows, row_findings = _report_rows(rows, report_type, presentation)
        export = {
            "file": path.name,
            "sha256": _file_sha256(path),
            "report_type": report_type,
            "closing_mode": closing_mode,
            "filename_scope": filename_scope,
            "header_scope": header_scope,
            "identity_ok": identity_ok,
            "date_ok": date_ok,
            "generation_timestamp": _generation_timestamp(rows),
            "required_for_g9_statement_parity": required,
            "visible_nonzero_row_count": sum(
                1 for row in normalized_rows if any(_decimal(row.get(field)) for field in AMOUNT_FIELDS)
            ),
            "findings": row_findings,
        }
        exports.append(export)
        issues = []
        if not identity_ok:
            issues.append("FILENAME_HEADER_SCOPE_MISMATCH")
        if not date_ok:
            issues.append("REPORT_DATE_MISMATCH")
        if row_findings:
            issues.append("ROW_NORMALIZATION_FINDINGS")
        if issues:
            destination = blockers if required else warnings
            destination.append({"file": path.name, "case": [report_type, closing_mode, filename_scope], "reasons": issues})
            continue
        if not required:
            continue
        key = (report_type, closing_mode, filename_scope)
        if key in cases:
            blockers.append({"reason": "DUPLICATE_REQUIRED_CASE", "case": list(key), "files": [cases[key]["source_file"], path.name]})
            continue
        cases[key] = {
            "report_type": report_type,
            "closing_mode": closing_mode,
            "scope": filename_scope,
            "source_file": path.name,
            "source_sha256": export["sha256"],
            "zero_row_policy": "source-omitted-rows-filled-from-frozen-presentation",
            "rows": normalized_rows,
            "controls": _presentation_controls(normalized_rows, presentation, report_type),
        }
    for key in sorted(REQUIRED_OFFICE_CASES - set(cases)):
        blockers.append({"reason": "REQUIRED_CASE_MISSING", "case": list(key)})
    normalized_cases = [cases[key] for key in sorted(cases)]
    if not blockers:
        normalized_cases.extend(_derived_consolidated(cases, presentation))
        normalized_cases.sort(key=lambda case: (case["report_type"], case["closing_mode"], case["scope"]))
    result = {
        "version": "arissto-december-2024-authoritative-statements-v1",
        "ready": not blockers,
        "period_id": "00053",
        "from_date": "2024-12-01",
        "to_date": "2024-12-31",
        "currency": "USD",
        "required_office_case_count": len(REQUIRED_OFFICE_CASES),
        "normalized_case_count": len(normalized_cases),
        "cases": normalized_cases,
        "source_exports": exports,
        "blockers": blockers,
        "warnings": warnings,
    }
    result["case_content_sha256"] = _content_sha256(result["cases"])
    return _json_value(result)


def target_presentation_controls(case: dict[str, Any], presentation: list[dict[str, Any]] | None = None) -> dict[str, Any]:
    return _json_value(_presentation_controls(case["rows"], presentation or _load_presentation(), case["report_type"]))
